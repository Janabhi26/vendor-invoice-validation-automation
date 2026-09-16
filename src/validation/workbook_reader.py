import os
from openpyxl import load_workbook


DEFAULT_WORKBOOK = os.path.abspath(
    os.path.join(
        os.path.dirname(__file__),
        "..",
        "..",
        "data",
        "Nassau.xlsx",
    )
)


WORKSHEETS = [
    "Offline Orders",
    "eBay, Amazon & Walmart",
    "NNC NES & Non-Wire",
    "Government Bidding",
    "Exports",
]


def normalize_value(value):
    """Convert a workbook value to a clean string."""
    if value is None:
        return None

    return str(value).strip()


def normalize_po(value):
    """Normalize PO values for reliable comparison."""
    if value is None:
        return None

    return str(value).strip().upper()


class NassauWorkbookReader:

    def __init__(self, workbook_path=DEFAULT_WORKBOOK):
        self.workbook_path = workbook_path
        self.workbook = None

    def open(self):
        """Open the Nassau workbook in read-only mode."""

        if not os.path.exists(self.workbook_path):
            raise FileNotFoundError(
                f"Workbook not found: {self.workbook_path}"
            )

        self.workbook = load_workbook(
            self.workbook_path,
            read_only=True,
            data_only=True,
        )

    def close(self):
        """Close the workbook."""

        if self.workbook is not None:
            self.workbook.close()
            self.workbook = None

    def _read_sheet(self, worksheet):
        """
        Read a worksheet using row 2 as the header row.

        The current Nassau workbook stores its column names
        on row 2 for the main order sheets.
        """

        rows = worksheet.iter_rows(
            min_row=2,
            values_only=True,
        )

        headers = next(rows, None)

        if not headers:
            return

        headers = [
            normalize_value(header)
            for header in headers
        ]

        for row in rows:

            if not any(value is not None for value in row):
                continue

            record = {}

            for index, header in enumerate(headers):

                if not header:
                    continue

                value = (
                    row[index]
                    if index < len(row)
                    else None
                )

                record[header] = value

            yield record

    def _build_result(
        self,
        sheet_name,
        matched_by,
        record,
    ):
        """
        Return the workbook fields needed by the
        invoice validation process.
        """

        return {
            "sheet": sheet_name,
            "matched_by": matched_by,
            "data": {
                "po": record.get("PO"),
                "product": record.get("Product"),
                "cost_per_m": record.get("Cost/M"),
                "cost": record.get("Cost"),
                "cut_charge": record.get("Cut Charge"),
                "freight": record.get("Freight"),
                "extended": record.get("Extended"),
                "extended_total_cost": record.get(
                    "Extended\n(Total Cost)"
                ),
                "carrier": record.get("Carrier"),
                "tracking_number": record.get(
                    "Tracking Number"
                ),
                "tracing_number": record.get(
                    "Tracing Number"
                ),
                "split_order_po": record.get(
                    "Split Order PO#"
                ),
                "qty": record.get("QTY"),
            },
            "raw_data": record,
        }

    def find_po(self, po_number):
        """
        Search for a PO in the configured Nassau sheets.

        First searches the PO column.
        Then searches Split Order PO#.
        """

        if self.workbook is None:
            self.open()

        target_po = normalize_po(po_number)

        if not target_po:
            return None

        # -------------------------------------------------
        # 1. Search the main PO column
        # -------------------------------------------------

        for sheet_name in WORKSHEETS:

            if sheet_name not in self.workbook.sheetnames:
                continue

            worksheet = self.workbook[sheet_name]

            for record in self._read_sheet(worksheet):

                record_po = normalize_po(
                    record.get("PO")
                )

                if record_po == target_po:

                    return self._build_result(
                        sheet_name,
                        "PO",
                        record,
                    )

        # -------------------------------------------------
        # 2. Search Split Order PO#
        # -------------------------------------------------

        for sheet_name in WORKSHEETS:

            if sheet_name not in self.workbook.sheetnames:
                continue

            worksheet = self.workbook[sheet_name]

            for record in self._read_sheet(worksheet):

                split_po = normalize_po(
                    record.get("Split Order PO#")
                )

                if split_po == target_po:

                    return self._build_result(
                        sheet_name,
                        "Split Order PO#",
                        record,
                    )

        return None


if __name__ == "__main__":

    reader = NassauWorkbookReader()

    try:

        print("=" * 80)
        print("NASSAU WORKBOOK TEST")
        print("=" * 80)

        print(
            "Workbook:",
            reader.workbook_path,
        )

        test_pos = [
            "E31276",
            "N159173",
        ]

        for test_po in test_pos:

            print()
            print("-" * 80)
            print(f"Searching for PO: {test_po}")
            print("-" * 80)

            result = reader.find_po(test_po)

            if result is None:

                print("❌ PO NOT FOUND")

                continue

            print("✅ PO FOUND")
            print(
                f"Sheet      : {result['sheet']}"
            )
            print(
                f"Matched By : {result['matched_by']}"
            )

            print()
            print("Fields for comparison:")

            for key, value in result["data"].items():

                print(
                    f"{key:25}: {value}"
                )

    finally:

        reader.close()
